"""A small Word file shaped like the owner's resume, for the docx tests:
header, headings with a rule under them, employer and title lines, bullets,
a column break and the two-column section. No personal content."""

import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _run(paragraph, text: str, bold: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold or None
    return run


def _border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
        bottom.set(qn(key), value)
    border.append(bottom)
    ppr.append(border)


def _numbered(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    for tag, value in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), value)
        num.append(el)
    ppr.append(num)


def _tabs(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "10620")
    tabs.append(tab)
    ppr.append(tabs)


def template() -> bytes:
    doc = Document()
    head = doc.add_paragraph(style="No Spacing")
    _run(head, "Sample Person", bold=True)
    _run(head, "        ")
    _run(head, "sample@example.com • 555-0100")
    sect = OxmlElement("w:sectPr")                      # the header is its own section
    head._p.get_or_add_pPr().append(sect)

    heading = doc.add_paragraph(style="No Spacing")
    _run(heading, "Experience", bold=True)
    _border(heading)
    doc.add_paragraph(style="No Spacing")               # spacer
    employer = doc.add_paragraph()
    _run(employer, "  Sample Employer", bold=True)
    title = doc.add_paragraph()
    _run(title, "    ", bold=True)
    _run(title, "Analyst – Somewhere, State            Jan 2020 - Present")
    _tabs(title)
    bullet = doc.add_paragraph(style="List Paragraph")
    _run(bullet, "Did a thing that mattered")
    _numbered(bullet)
    brk = doc.add_paragraph(style="No Spacing")
    brk.add_run().add_break(WD_BREAK.COLUMN)
    school = doc.add_paragraph(style="No Spacing")
    _run(school, "  ", bold=True)
    _run(school, "Sample University ", bold=True)
    _run(school, "| May 2020")
    degree = doc.add_paragraph(style="No Spacing")
    _run(degree, "    B.S. in Something | 4.00/4.00")
    pbullet = doc.add_paragraph(style="No Spacing")
    _run(pbullet, "A project bullet")
    _numbered(pbullet)

    body = doc.element.body
    end = body.find(qn("w:sectPr"))
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:equalWidth"), "0")
    for width in ("6840", "4392"):
        col = OxmlElement("w:col")
        col.set(qn("w:w"), width)
        cols.append(col)
    end.append(cols)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


SAMPLE = {
    "name": "Sample Person", "email": "sample@example.com", "phone": "555-0100",
    "linkedin": "linkedin.com/in/sample",
    "roles": [{"employer": "Acme Energy", "title": "Business Analyst", "location": "Denver, CO",
               "start": "Jan 2025", "end": "Present", "bullets": [
                   {"text": "Built a pipeline that saved 2 hours a day.", "skills": ["Python", "SQL"],
                    "source": "resume", "rank": 4, "in_master": True},
                   {"text": "Pool bullet, not on the master.", "skills": [], "source": "history",
                    "rank": 2, "in_master": False}]},
              {"employer": "Beta Chemicals", "title": "Project Engineer", "location": "Houston, TX",
               "start": "Jun 2023", "end": "Dec 2024", "bullets": [
                   {"text": "Led projects.", "skills": [], "source": "resume", "rank": 3,
                    "in_master": True}]}],
    "projects": [{"employer": "Workout Analyzer", "title": "", "location": "", "start": "",
                  "end": "", "bullets": [{"text": "Developed a tool.", "skills": ["VBA"],
                                          "source": "resume", "rank": 3, "in_master": True}]}],
    "education": [{"school": "Sample University", "when": "May 2026",
                   "degree": "M.S. in Data Science | 3.90/4.00"}],
    "skill_groups": [{"name": "Programming Languages", "skills": "Python, SQL, DAX"}],
    "conflicts": [],
}
